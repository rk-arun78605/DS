import streamlit as st
import psycopg2
from datetime import datetime, timedelta, date
from psycopg2.extras import RealDictCursor, execute_values
import html
import plotly.graph_objects as go
import streamlit.components.v1 as components
import io
import base64
import os
import re
import warnings

import pandas as pd

try:
    import wms_sync as _wms_sync
    WMS_SYNC_AVAILABLE = True
except Exception:
    WMS_SYNC_AVAILABLE = False

# ── LVO upload column list (matches home_dashboard TABLE_CONFIGS) ──────────────
_LVO_COLUMNS = [
    "date", "shop_code", "vehicle_no", "item_code", "item_name",
    "qty_loaded", "value_loaded", "qty_offloaded", "value_offloaded",
    "diff_qty", "diff_val", "price", "diff",
]

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    OPENPYXL_AVAILABLE = True
except Exception:
    Workbook = None
    PatternFill = None
    Font = None
    Alignment = None
    OPENPYXL_AVAILABLE = False

st.set_page_config(
    page_title="Loading vs Offloading Dashboard",
    page_icon="https://melcom.com/media/favicon/stores/1/faviconn_162_x_184px.jpg",
    layout="wide",
    initial_sidebar_state="collapsed",
)

DB_CONFIG = {
    "host": "localhost",
    "port": 3307,
    "user": "postgres",
    "password": "hello",
    "dbname": "WH",
}

USERS_DB_NAME = "users"
CENTURY_DB_NAME = "century_penetration"
REQUIRED_TABLE_ACCESS = "offloading_vs_loading"


def inject_css() -> None:
    st.markdown(
        """
        <style>
        #MainMenu, footer, header { visibility: hidden; }
        html, body, .stApp {
            background: #0a0f1e !important;
            color: #e2e8f0 !important;
            font-family: 'Inter', 'Segoe UI', sans-serif !important;
        }
        .block-container {
            padding-top: 0.7rem !important;
            padding-left: 1rem !important;
            padding-right: 1rem !important;
            max-width: 100% !important;
        }
        .dashboard-title {
            color: #f1f5f9;
            font-size: 1.3rem;
            font-weight: 800;
            line-height: 1.15;
        }
        .dashboard-sub {
            color: #94a3b8;
            font-size: 0.78rem;
            font-weight: 500;
        }
        .section-title {
            color: #dbeafe;
            font-size: 1rem;
            font-weight: 700;
            margin: 0.5rem 0 0.5rem;
        }
        div[data-testid="stTextInput"] label p,
        div[data-testid="stSelectbox"] label p,
        div[data-testid="stPasswordInput"] label p,
        div[data-testid="stTextInput"] input,
        div[data-testid="stPasswordInput"] input {
            color: #f8fafc !important;
        }
        div[data-testid="stTextInput"] input,
        div[data-testid="stPasswordInput"] input {
            background: rgba(15, 23, 42, 0.8) !important;
            border: 1px solid rgba(59, 130, 246, 0.35) !important;
        }
        div.stButton > button,
        div[data-testid="stFormSubmitButton"] > button {
            color: #ffffff !important;
            background: linear-gradient(135deg, #1d4ed8 0%, #3b82f6 100%) !important;
            border: 1px solid rgba(191, 219, 254, 0.65) !important;
            font-weight: 700 !important;
        }
        /* Reset Password popover toggle button — dark navy to match dashboard */
        button[data-testid="stPopoverButton"] {
            color: #ffffff !important;
            background: linear-gradient(135deg, #0c1a3a 0%, #112244 100%) !important;
            border: 1px solid rgba(59, 130, 246, 0.40) !important;
            font-weight: 600 !important;
        }
        button[data-testid="stPopoverButton"]:hover {
            background: linear-gradient(135deg, #1a2f5a 0%, #1e3872 100%) !important;
            border-color: rgba(96, 165, 250, 0.60) !important;
        }

        /* ── Popover dropdown panel — dark dashboard theme ── */
        div[data-testid="stPopover"] > div,
        div[data-testid="stPopoverContent"],
        [data-baseweb="popover"] [role="dialog"],
        [data-baseweb="popover"] > div {
            background: linear-gradient(160deg, #0c1a3a 0%, #080f22 100%) !important;
            border: 1px solid rgba(59, 130, 246, 0.30) !important;
            border-radius: 10px !important;
        }
        /* All text inside the popover panel */
        div[data-testid="stPopoverContent"] *,
        [data-baseweb="popover"] [role="dialog"] * {
            color: #f8fafc !important;
        }
        /* Input fields inside popover */
        div[data-testid="stPopoverContent"] input,
        [data-baseweb="popover"] [role="dialog"] input {
            background: rgba(15, 23, 42, 0.85) !important;
            border: 1px solid rgba(59, 130, 246, 0.35) !important;
            color: #f8fafc !important;
        }
        /* Labels inside popover */
        div[data-testid="stPopoverContent"] label p,
        [data-baseweb="popover"] [role="dialog"] label p {
            color: #cbd5e1 !important;
        }
        /* Submit button inside popover */
        div[data-testid="stPopoverContent"] button,
        [data-baseweb="popover"] [role="dialog"] button {
            color: #ffffff !important;
            background: linear-gradient(135deg, #1d4ed8 0%, #3b82f6 100%) !important;
            border: 1px solid rgba(191, 219, 254, 0.55) !important;
            font-weight: 700 !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def get_db_connection(dbname: str | None = None):
    cfg = dict(DB_CONFIG)
    if dbname:
        cfg["dbname"] = dbname
    return psycopg2.connect(**cfg)


@st.cache_data(ttl=300, show_spinner=False)
def _wms_check_file_status(target_date: date) -> str:
    """
    Returns:
      'ok'            — file exists on network share
      'missing'       — share accessible but file not there
      'network_error' — share not reachable
    Cached for 5 minutes so it doesn't hammer the network share on every rerun.
    """
    if not WMS_SYNC_AVAILABLE:
        return "ok"
    try:
        found = _wms_sync.find_wms_file(target_date)
        return "ok" if found else "missing"
    except OSError:
        return "network_error"


def init_auth_state() -> None:
    if "ovl_authenticated" not in st.session_state:
        st.session_state["ovl_authenticated"] = False
    if "ovl_user" not in st.session_state:
        st.session_state["ovl_user"] = None
    if "ovl_is_master_user" not in st.session_state:
        st.session_state["ovl_is_master_user"] = False
    if "ovl_allowed_shops" not in st.session_state:
        st.session_state["ovl_allowed_shops"] = []


@st.cache_data(ttl=300, show_spinner=False)
def authenticate_user(login_name: str, password: str) -> dict | None:
    query = """
        SELECT employee_id, full_name, table_access, is_active
        FROM users
        WHERE (
                UPPER(TRIM(COALESCE(full_name, ''))) = UPPER(TRIM(%s))
             OR UPPER(TRIM(COALESCE(employee_id, ''))) = UPPER(TRIM(%s))
        )
          AND password = %s
          AND LOWER(COALESCE(is_active, 'false')) = 'true'
        LIMIT 1
    """
    with get_db_connection(USERS_DB_NAME) as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(query, (login_name.strip(), login_name.strip(), password))
            row = cur.fetchone()
    return dict(row) if row else None


def check_dashboard_access(user: dict) -> bool:
    table_access = str((user or {}).get("table_access", "") or "")
    allowed = {x.strip().lower() for x in table_access.split(",") if x.strip()}
    return bool(
        "all" in allowed
        or REQUIRED_TABLE_ACCESS in allowed
        or "offloading_loading" in allowed
        or "offloading_vs_loading_dashboard" in allowed
    )


def get_user_shop_scope(employee_id: str, full_name: str) -> tuple[bool, list[str]]:
    emp = str(employee_id or "").strip().upper()
    name = str(full_name or "").strip().upper()
    query = """
        WITH latest_upload AS (
            SELECT MAX(upload_date) AS max_upload_date
            FROM opsmgr
        )
        SELECT UPPER(TRIM(COALESCE(shopcode, ''))) AS shop_code
        FROM opsmgr
        WHERE upload_date = (SELECT max_upload_date FROM latest_upload)
          AND (SELECT max_upload_date FROM latest_upload) IS NOT NULL
          AND UPPER(TRIM(COALESCE(opsmanagername, ''))) IN (%s, %s)
          AND TRIM(COALESCE(shopcode, '')) <> ''
    """

    with get_db_connection(CENTURY_DB_NAME) as conn:
        with conn.cursor() as cur:
            cur.execute(query, (emp, name))
            rows = cur.fetchall()

    shops = sorted({str(r[0] or "").strip().upper() for r in rows if r and str(r[0] or "").strip()})
    is_master = any(s == "ALL" for s in shops)
    scoped = [s for s in shops if s != "ALL"]
    return is_master, scoped


def get_latest_opsmgr_upload_date() -> date | None:
    query = """
        SELECT MAX(upload_date)
        FROM opsmgr
    """
    with get_db_connection(CENTURY_DB_NAME) as conn:
        with conn.cursor() as cur:
            cur.execute(query)
            row = cur.fetchone()
    if not row:
        return None
    return row[0]


def change_password(login_name: str, new_password: str) -> tuple[bool, str]:
    if not login_name.strip() or not new_password:
        return False, "Please fill all password fields."
    if len(new_password) < 4:
        return False, "New password must be at least 4 characters."
    resolve_q = """
                SELECT employee_id
        FROM users
                WHERE (
                                UPPER(TRIM(COALESCE(full_name, ''))) = UPPER(TRIM(%s))
                         OR UPPER(TRIM(COALESCE(employee_id, ''))) = UPPER(TRIM(%s))
                )
          AND LOWER(COALESCE(is_active, 'false')) = 'true'
        LIMIT 1
    """
    update_q = """
        UPDATE users
        SET password = %s
        WHERE employee_id = %s
    """
    with get_db_connection(USERS_DB_NAME) as conn:
        with conn.cursor() as cur:
            cur.execute(resolve_q, (login_name.strip(), login_name.strip()))
            ok_row = cur.fetchone()
            if not ok_row:
                return False, "Invalid User Name / User ID or inactive account."
            target_employee_id = str(ok_row[0] or "").strip()
            if not target_employee_id:
                return False, "Unable to resolve user account for password reset."
            cur.execute(update_q, (new_password, target_employee_id))
        conn.commit()

    authenticate_user.clear()
    return True, "Password updated successfully."


def render_login_gate() -> bool:
    st.markdown(
        """
        <div style='max-width:520px;margin:28px auto 0 auto;padding:16px;border:1px solid rgba(59,130,246,0.25);border-radius:12px;background:linear-gradient(160deg,#0c1a3a 0%,#080f22 100%);'>
            <div style='display:flex;align-items:center;gap:10px;margin-bottom:8px;'>
                <img src='https://melcom.com/media/favicon/stores/1/faviconn_162_x_184px.jpg' style='height:38px;width:38px;border-radius:50%;border:1px solid rgba(148,163,184,0.45);' />
                <div style='font-size:1.1rem;font-weight:800;color:#f8fafc;'>Offloading vs Loading - Login</div>
            </div>
            <div style='font-size:0.82rem;color:#cbd5e1;margin-bottom:2px;'>Sign in with User Name and password.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    left_col, center_col, right_col = st.columns([1.3, 1.5, 1.3], gap="small")
    with center_col:
        with st.form("ovl_login_form", clear_on_submit=False):
            ops_manager_name = st.text_input("User Name")
            password = st.text_input("Password", type="password")
            submitted = st.form_submit_button("Login", use_container_width=True)

    if submitted:
        user = authenticate_user(ops_manager_name, password)
        if not user:
            st.error("Invalid login credentials or inactive account.")
            return False
        if not check_dashboard_access(user):
            st.error("You do not have access to this dashboard. Ask admin to update users.table_access.")
            return False

        # Super user: table_access='all' → auto master, no opsmgr entry required
        _ta = str((user or {}).get("table_access", "") or "")
        _is_super = "all" in {x.strip().lower() for x in _ta.split(",") if x.strip()}
        if _is_super:
            is_master, allowed_shops = True, []
        else:
            is_master, allowed_shops = get_user_shop_scope(
                str(user.get("employee_id", "") or ""),
                str(user.get("full_name", "") or ""),
            )
        if not _is_super and not is_master and not allowed_shops:
            st.error("No shop mapping found in century_penetration.opsmgr for this user.")
            return False

        st.session_state["ovl_authenticated"] = True
        st.session_state["ovl_user"] = user
        st.session_state["ovl_is_master_user"] = is_master
        st.session_state["ovl_allowed_shops"] = allowed_shops
        st.rerun()

    with center_col:
        with st.popover("Reset Password"):
            with st.form("ovl_reset_pwd_login_form", clear_on_submit=True):
                emp = st.text_input("User Name", key="ovl_reset_emp")
                new_pwd = st.text_input("New Password", type="password", key="ovl_reset_new")
                cnf_pwd = st.text_input("Confirm New Password", type="password", key="ovl_reset_confirm")
                do_reset = st.form_submit_button("Update Password", use_container_width=True)

            if do_reset:
                if new_pwd != cnf_pwd:
                    st.error("New password and confirm password do not match.")
                else:
                    ok, msg = change_password(emp, new_pwd)
                    if ok:
                        st.success(msg)
                    else:
                        st.error(msg)

    return bool(st.session_state.get("ovl_authenticated", False))


def render_logged_in_sidebar() -> bool:
    user = st.session_state.get("ovl_user") or {}
    full_name = str(user.get("full_name") or "").strip()
    employee_id = str(user.get("employee_id") or "").strip()
    is_master = bool(st.session_state.get("ovl_is_master_user", False))
    allowed_shops = st.session_state.get("ovl_allowed_shops", []) or []

    with st.sidebar:
        st.markdown("### User Access")
        st.caption(f"Name: {full_name or 'Unknown'}")
        st.caption(f"User ID: {employee_id or 'Unknown'}")
        if is_master:
            st.success("Role: Master User (ALL shops)")
        else:
            st.info(f"Role: Ops Manager ({len(allowed_shops)} mapped shop(s))")
            if allowed_shops:
                st.caption("Shops: " + ", ".join(allowed_shops))

        with st.expander("Reset Password", expanded=False):
            with st.form("ovl_reset_pwd_loggedin_form", clear_on_submit=True):
                login_name = st.text_input("User Name", key="ovl_lg_reset_name")
                new_pwd = st.text_input("New Password", type="password", key="ovl_lg_reset_new")
                cnf_pwd = st.text_input("Confirm New Password", type="password", key="ovl_lg_reset_confirm")
                do_reset = st.form_submit_button("Update Password", use_container_width=True)

            if do_reset:
                if new_pwd != cnf_pwd:
                    st.error("New password and confirm password do not match.")
                else:
                    reset_name = str(login_name or "").strip() or full_name or employee_id
                    ok, msg = change_password(reset_name, new_pwd)
                    if ok:
                        st.success(msg)
                    else:
                        st.error(msg)

        if st.button("Logout", use_container_width=True):
            st.session_state["ovl_authenticated"] = False
            st.session_state["ovl_user"] = None
            st.session_state["ovl_is_master_user"] = False
            st.session_state["ovl_allowed_shops"] = []
            st.rerun()
            return True

    return False


def apply_shop_scope(rows: list[dict], allowed_shops: list[str], is_master: bool) -> list[dict]:
    if is_master:
        return rows
    allowed = {str(s or "").strip().upper() for s in (allowed_shops or []) if str(s or "").strip()}
    if not allowed:
        return []
    return [
        r for r in rows
        if str(r.get("shop_code", "") or "").strip().upper() in allowed
    ]


def corrected_diff_val_sql() -> str:
    """Return SQL expression that normalizes legacy diff_val sign issues."""
    return """
        CASE
            WHEN COALESCE(diff_val, 0) = 0
                 AND COALESCE(diff_qty, 0) <> 0
                 AND COALESCE(price, 0) <> 0
                THEN COALESCE(diff_qty, 0) * COALESCE(price, 0)
            WHEN COALESCE(diff_qty, 0) < 0
                 AND COALESCE(diff_val, 0) > 0
                 AND COALESCE(price, 0) > 0
                THEN COALESCE(diff_qty, 0) * COALESCE(price, 0)
            WHEN COALESCE(diff_qty, 0) > 0
                 AND COALESCE(diff_val, 0) < 0
                 AND COALESCE(price, 0) > 0
                THEN COALESCE(diff_qty, 0) * COALESCE(price, 0)
            ELSE COALESCE(diff_val, 0)
        END
    """


def get_live_date_bounds_uncached() -> tuple:
    """Read bounds directly from DB to recover from stale cache states."""
    query = """
        SELECT COUNT(*) AS row_count, MIN(date) AS min_date, MAX(date) AS max_date
        FROM offloading_vs_loading
        WHERE date IS NOT NULL
    """
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(query)
            row = cur.fetchone()
    if not row:
        return 0, None, None
    return int(row[0] or 0), row[1], row[2]


@st.cache_data(ttl=300)
def load_date_bounds() -> tuple:
    query = """
        SELECT MIN(date) AS min_date, MAX(date) AS max_date
        FROM offloading_vs_loading
        WHERE date IS NOT NULL
    """
    with get_db_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(query)
            row = cur.fetchone()
    if not row or row[0] is None or row[1] is None:
        return None, None
    return row[0], row[1]


@st.cache_data(ttl=300)
def load_offloading_vs_loading(start_date, end_date) -> list[dict]:
    diff_val_expr = "((COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) * COALESCE(price, 0))"
    query = f"""
        SELECT
            date,
            shop_code,
            vehicle_no,
            item_code,
            item_name,
            COALESCE(qty_loaded, 0)::numeric AS qty_loaded,
            (COALESCE(qty_loaded, 0) * COALESCE(price, 0))::numeric AS value_loaded,
            COALESCE(qty_offloaded, 0)::numeric AS qty_offloaded,
            (COALESCE(qty_offloaded, 0) * COALESCE(price, 0))::numeric AS value_offloaded,
            (COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0))::numeric AS diff_qty,
            ({diff_val_expr})::numeric AS diff_val
        FROM offloading_vs_loading
        WHERE date BETWEEN %(s)s AND %(e)s
        ORDER BY date DESC, shop_code, vehicle_no, item_code
    """
    with get_db_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(query, {"s": start_date, "e": end_date})
            return [dict(r) for r in cur.fetchall()]


@st.cache_data(ttl=300)
def load_monthly_summary(start_date, end_date) -> list[dict]:
    diff_val_expr = "((COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) * COALESCE(price, 0))"
    query = f"""
        SELECT
            to_char(date_trunc('month', date), 'YYYY-MM') AS month_key,
            to_char(date_trunc('month', date), 'Mon YYYY') AS month_label,
            SUM(COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0))::numeric AS diff_qty,
            SUM({diff_val_expr})::numeric AS diff_val
        FROM offloading_vs_loading
        WHERE date BETWEEN %(s)s AND %(e)s
        GROUP BY 1, 2
        ORDER BY month_key
    """
    with get_db_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(query, {"s": start_date, "e": end_date})
            return [dict(r) for r in cur.fetchall()]


@st.cache_data(ttl=300)
def load_monthly_summary_all() -> list[dict]:
    diff_val_expr = "((COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) * COALESCE(price, 0))"
    query = f"""
        SELECT
            to_char(date_trunc('month', date), 'YYYY-MM') AS month_key,
            to_char(date_trunc('month', date), 'Mon YYYY') AS month_label,
            SUM(COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0))::numeric AS diff_qty,
            SUM({diff_val_expr})::numeric AS diff_val
        FROM offloading_vs_loading
        WHERE date IS NOT NULL
        GROUP BY 1, 2
        ORDER BY month_key
    """
    with get_db_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(query)
            return [dict(r) for r in cur.fetchall()]


@st.cache_data(ttl=300)
def load_shop_summary(start_date, end_date) -> list[dict]:
    diff_val_expr = "((COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) * COALESCE(price, 0))"
    query = f"""
        SELECT
            shop_code,
            SUM(COALESCE(qty_loaded, 0))::numeric AS qty_loaded,
            SUM(COALESCE(qty_loaded, 0) * COALESCE(price, 0))::numeric AS value_loaded,
            SUM(COALESCE(qty_offloaded, 0))::numeric AS qty_offloaded,
            SUM(COALESCE(qty_offloaded, 0) * COALESCE(price, 0))::numeric AS value_offloaded,
            SUM(CASE WHEN (COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) < 0 THEN ABS(COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) ELSE 0 END)::numeric AS short_qty,
            SUM(CASE WHEN ({diff_val_expr}) < 0 THEN ABS({diff_val_expr}) ELSE 0 END)::numeric AS short_val,
            SUM(CASE WHEN (COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) > 0 THEN (COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0)) ELSE 0 END)::numeric AS excess_qty,
            SUM(CASE WHEN ({diff_val_expr}) > 0 THEN ({diff_val_expr}) ELSE 0 END)::numeric AS excess_val,
            SUM(COALESCE(qty_offloaded, 0) - COALESCE(qty_loaded, 0))::numeric AS diff_qty,
            SUM({diff_val_expr})::numeric AS diff_val
        FROM offloading_vs_loading
        WHERE date BETWEEN %(s)s AND %(e)s
        GROUP BY shop_code
        ORDER BY ABS(SUM({diff_val_expr})) DESC, ABS(SUM(diff_qty)) DESC
    """
    with get_db_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(query, {"s": start_date, "e": end_date})
            return [dict(r) for r in cur.fetchall()]


@st.cache_data(ttl=300)
def load_shop_meta_map() -> dict[str, dict[str, str]]:
    query = """
        WITH mgr AS (
            SELECT
                UPPER(TRIM(shop_code)) AS shop_code,
                MAX(TRIM(shop_description)) AS shop_name,
                STRING_AGG(DISTINCT TRIM(shop_manager_name), ', ' ORDER BY TRIM(shop_manager_name)) AS manager_name
            FROM shopmgrname
            WHERE is_current = TRUE
              AND TRIM(COALESCE(shop_code, '')) <> ''
            GROUP BY UPPER(TRIM(shop_code))
        ),
        stg AS (
            SELECT
                UPPER(TRIM(shop_code)) AS shop_code,
                MAX(TRIM(shop_name)) AS shop_name
            FROM offloading_loading_staging
            WHERE TRIM(COALESCE(shop_code, '')) <> ''
            GROUP BY UPPER(TRIM(shop_code))
        )
        SELECT
            COALESCE(mgr.shop_code, stg.shop_code) AS shop_code,
            COALESCE(NULLIF(mgr.shop_name, ''), stg.shop_name, '') AS shop_name,
            COALESCE(mgr.manager_name, '') AS manager_name
        FROM stg
        FULL OUTER JOIN mgr ON mgr.shop_code = stg.shop_code
    """
    with get_db_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(query)
            rows = cur.fetchall()

    result: dict[str, dict[str, str]] = {}
    for r in rows:
        code = str(r.get("shop_code", "") or "").strip().upper()
        if not code:
            continue
        result[code] = {
            "shop_name": str(r.get("shop_name", "") or "").strip(),
            "manager_name": str(r.get("manager_name", "") or "").strip(),
        }
    return result


def render_metrics(rows: list[dict], start_date, end_date) -> None:
    total_wh_loaded_qty = sum(float(r.get("qty_loaded") or 0) for r in rows)
    total_shop_receiving_qty = sum(float(r.get("qty_offloaded") or 0) for r in rows)
    total_diff_qty = sum(float(r.get("diff_qty") or 0) for r in rows)
    total_diff_val = sum(float(r.get("diff_val") or 0) for r in rows)

    st.markdown(
        f"<div style='color:#94a3b8;font-size:0.8rem;margin:0.2rem 0 0.6rem;'>Selected Range: {start_date.strftime('%d %b %Y')} → {end_date.strftime('%d %b %Y')}</div>",
        unsafe_allow_html=True,
    )

    cols = st.columns(4)
    cards = [
        ("WH Loaded Qty", f"{total_wh_loaded_qty:,.0f}"),
        ("Shop Receiving Qty", f"{total_shop_receiving_qty:,.0f}"),
        ("Final Difference QTY", f"{total_diff_qty:,.0f}"),
        ("Final Difference Value", f"{total_diff_val:,.2f}"),
    ]

    for col, (label, value) in zip(cols, cards):
        with col:
            st.markdown(
                f"""
                <div style="background:linear-gradient(160deg,#0c1a3a 0%,#080f22 100%);border:1px solid rgba(59,130,246,0.22);border-radius:14px;padding:12px;">
                    <div style="color:#94a3b8;font-size:0.70rem;text-transform:uppercase;font-weight:700;letter-spacing:0.05em;">{label}</div>
                    <div style="color:#f8fafc;font-size:1.4rem;font-weight:800;line-height:1.15;margin-top:2px;">{value}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )


def render_monthly_graph(monthly_rows: list[dict]) -> None:
    if not monthly_rows:
        st.info("No monthly data for selected range.")
        return

    x_vals = [r.get("month_label", "") for r in monthly_rows]
    diff_val = [float(r.get("diff_val") or 0) for r in monthly_rows]

    fig = go.Figure()
    fig.add_trace(
        go.Scatter(
            x=x_vals,
            y=diff_val,
            mode="lines+markers",
            name="Final Difference Value",
            line=dict(color="#f59e0b", width=3),
            marker=dict(size=7, color="#f59e0b"),
        )
    )

    fig.update_layout(
        title="Month-on-Month Difference Value",
        xaxis_title="Month",
        yaxis_title="Final Difference Value",
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(17,26,51,0.65)",
        font=dict(color="#e2e8f0"),
        margin=dict(l=30, r=30, t=50, b=30),
        height=340,
    )
    st.plotly_chart(fig, use_container_width=True)


def render_sortable_html_table(
    title: str,
    headers: list[str],
    rows: list[list[str]],
    sort_values: list[list[str]],
    table_id: str,
    min_width_px: int = 1120,
    row_link_values: list[str] | None = None,
    row_link_param: str = "detail_shop",
    show_title: bool = True,
    cell_styles: dict[tuple[int, int], str] | None = None,
    enable_download_hover: bool = False,
    download_file_name: str = "table_export.csv",
    excel_base64: str | None = None,
    total_row: list[str] | None = None,
) -> None:
    if not rows:
        st.info("No rows to display.")
        return

    header_html = "".join(f"<th data-col='{idx}'>{html.escape(h)} ⬍</th>" for idx, h in enumerate(headers))
    body_parts = []
    for row_idx, row in enumerate(rows):
        cells = []
        row_shop = None
        if row_link_values and row_idx < len(row_link_values):
            link_value = str(row_link_values[row_idx] or "").strip()
            if link_value:
                row_shop = link_value
        for col_idx, val in enumerate(row):
            sval = sort_values[row_idx][col_idx]
            style_attr = ""
            if cell_styles:
                style_txt = cell_styles.get((row_idx, col_idx), "")
                if style_txt:
                    style_attr = f" style='{style_txt}'"
            if row_shop:
                cells.append(
                    f"<td data-sort='{html.escape(str(sval))}' data-shop='{html.escape(row_shop)}' class='shop-click-cell' {style_attr}>{html.escape(str(val))}</td>"
                )
            else:
                cells.append(f"<td data-sort='{html.escape(str(sval))}'{style_attr}>{html.escape(str(val))}</td>")
        body_parts.append("<tr>" + "".join(cells) + "</tr>")
    body_html = "".join(body_parts)

    if show_title:
        st.markdown(
            f"<div class='section-title'>{title}</div>",
            unsafe_allow_html=True,
        )

    download_button_html = ""
    if enable_download_hover:
        download_button_html = f"<button id='{table_id}-download' class='table-download-btn' title='Download CSV'>⭳</button>"

    tfoot_html = ""
    if total_row is not None:
        tfoot_cells = "".join(f"<td>{html.escape(str(v))}</td>" for v in total_row)
        tfoot_html = f"<tfoot><tr>{tfoot_cells}</tr></tfoot>"

    table_html = f"""
        <style>
        .table-shell-{table_id} {{
            position: relative;
        }}
        .table-download-btn {{
            position: absolute;
            top: 8px;
            right: 10px;
            z-index: 20;
            border: 1px solid rgba(96, 165, 250, 0.95);
            background: linear-gradient(135deg, #2563eb 0%, #3b82f6 100%);
            color: #ffffff;
            border-radius: 8px;
            font-size: 13px;
            line-height: 1;
            padding: 7px 9px;
            cursor: pointer;
            opacity: 0.92;
            transition: opacity .18s ease;
            box-shadow: 0 0 0 2px rgba(96, 165, 250, 0.28), 0 8px 20px rgba(37, 99, 235, 0.45);
        }}
        .table-shell-{table_id}:hover .table-download-btn {{
            opacity: 1;
        }}
        .table-download-btn:hover {{
            background: linear-gradient(135deg, #1d4ed8 0%, #2563eb 100%);
            transform: translateY(-1px);
            box-shadow: 0 0 0 2px rgba(147, 197, 253, 0.40), 0 10px 24px rgba(29, 78, 216, 0.55);
        }}
        #{table_id}-wrap {{
            border: 1px solid rgba(139, 92, 246, 0.3);
            border-radius: 14px;
            overflow: auto;
            max-height: 620px;
            background: linear-gradient(135deg, #111a33 0%, #0b1226 100%);
        }}
        #{table_id} {{
            width: 100%;
            border-collapse: collapse;
            min-width: {min_width_px}px;
            font-family: Inter, Segoe UI, sans-serif;
        }}
        #{table_id} th {{
            background: rgba(30, 58, 138, 0.55);
            color: #e2e8f0;
            font-size: 0.76rem;
            text-transform: uppercase;
            letter-spacing: 0.03em;
            text-align: center;
            padding: 8px;
            border-bottom: 1px solid rgba(139, 92, 246, 0.35);
            position: sticky;
            top: 0;
            z-index: 2;
            white-space: nowrap;
            cursor: pointer;
            user-select: none;
        }}
        #{table_id} td {{
            color: #e2e8f0;
            font-size: 0.8rem;
            text-align: center;
            padding: 7px;
            border-bottom: 1px solid rgba(71, 85, 105, 0.28);
            white-space: nowrap;
        }}
        #{table_id} tbody tr:hover {{
            background: rgba(99, 102, 241, 0.12);
        }}
        #{table_id} td:nth-child(6) {{
            text-align: center;
            max-width: 420px;
            overflow: hidden;
            text-overflow: ellipsis;
        }}
        #{table_id} td a:hover {{
            text-decoration: underline;
        }}
        #{table_id} td.shop-click-cell {{
            cursor: pointer;
        }}
        #{table_id} td.shop-click-cell:hover {{
            background: rgba(59, 130, 246, 0.18);
        }}
        #{table_id} tfoot td {{
            background: rgba(15, 30, 80, 0.95);
            color: #f8fafc;
            font-weight: 800;
            font-size: 0.8rem;
            text-align: center;
            border-top: 2px solid rgba(139, 92, 246, 0.6);
            position: sticky;
            bottom: 0;
            z-index: 1;
            white-space: nowrap;
        }}
        </style>
        <div class="table-shell-{table_id}">
            {download_button_html}
            <div id="{table_id}-wrap">
                <table id="{table_id}">
                    <thead><tr>{header_html}</tr></thead>
                    <tbody>{body_html}</tbody>
                    {tfoot_html}
                </table>
            </div>
        </div>
        <script>
        (() => {{
            const table = document.getElementById('{table_id}');
            if (!table) return;
            const headers = table.querySelectorAll('thead th');
            const tbody = table.querySelector('tbody');
            let sortState = {{ col: -1, asc: true }};

            const parseVal = (raw) => {{
                const cleaned = String(raw || '').replace(/,/g, '').trim();
                const n = Number(cleaned);
                return Number.isNaN(n) ? cleaned.toLowerCase() : n;
            }};

            headers.forEach((th, idx) => {{
                th.addEventListener('click', () => {{
                    const rows = Array.from(tbody.querySelectorAll('tr'));
                    const asc = sortState.col === idx ? !sortState.asc : true;
                    rows.sort((a, b) => {{
                        const av = parseVal(a.children[idx].getAttribute('data-sort'));
                        const bv = parseVal(b.children[idx].getAttribute('data-sort'));
                        if (av < bv) return asc ? -1 : 1;
                        if (av > bv) return asc ? 1 : -1;
                        return 0;
                    }});
                    rows.forEach(r => tbody.appendChild(r));
                    sortState = {{ col: idx, asc }};
                }});
            }});

            const clickableCells = table.querySelectorAll('td.shop-click-cell[data-shop]');
            clickableCells.forEach((td) => {{
                td.addEventListener('click', () => {{
                    const shop = String(td.getAttribute('data-shop') || '').trim();
                    if (!shop) return;
                    try {{
                        const baseUrl = new URL(window.parent.location.href);
                        baseUrl.searchParams.set('{row_link_param}', shop);
                        window.parent.location.href = baseUrl.toString();
                    }} catch (e) {{
                        const fallback = `?{row_link_param}=${{encodeURIComponent(shop)}}`;
                        window.parent.location.href = fallback;
                    }}
                }});
            }});

            const downloadBtn = document.getElementById('{table_id}-download');
            if (downloadBtn) {{
                downloadBtn.addEventListener('click', () => {{
                    const link = document.createElement('a');
                    const excelB64 = '{excel_base64 or ""}';
                    if (excelB64) {{
                        link.href = `data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,${{excelB64}}`;
                    }} else {{
                        const csvEscape = (v) => `"${{String(v ?? '').replace(/"/g, '""')}}"`;
                        const headerVals = Array.from(table.querySelectorAll('thead th')).map(th => th.innerText.replace('⬍', '').trim());
                        const bodyVals = Array.from(table.querySelectorAll('tbody tr')).map(tr =>
                            Array.from(tr.querySelectorAll('td')).map(td => td.innerText.trim())
                        );
                        const csvLines = [headerVals, ...bodyVals].map(row => row.map(csvEscape).join(','));
                        const csv = csvLines.join('\\n');
                        const blob = new Blob([csv], {{ type: 'text/csv;charset=utf-8;' }});
                        link.href = URL.createObjectURL(blob);
                    }}
                    link.download = '{download_file_name}';
                    document.body.appendChild(link);
                    link.click();
                    document.body.removeChild(link);
                    if (!excelB64) URL.revokeObjectURL(link.href);
                }});
            }}
        }})();
        </script>
        """
    extra = 40 if total_row is not None else 0
    viewport_height = min(640, max(240, 86 + len(rows) * 30 + extra))
    components.html(table_html, height=viewport_height, scrolling=False)


def render_detail_table(
    rows: list[dict],
    shop_meta_map: dict[str, dict[str, str]],
    export_rows: list[dict] | None = None,
) -> None:
    headers = [
        "Date",
        "Shop Manager",
        "Shop Name",
        "Vehicle No",
        "Item Code",
        "Item Name",
        "WH Loaded Qty",
        "WH Loaded Value",
        "Shop Receiving Qty",
        "Shop Receiving Value",
        "Final Difference QTY",
        "Final Difference Value",
    ]

    def num_fmt(v, decimals=2):
        try:
            return f"{float(v):,.{decimals}f}"
        except Exception:
            return "0.00" if decimals == 2 else "0"

    table_rows = []
    sort_rows = []

    def diff_cell_style(raw_value: float) -> str:
        try:
            num = float(raw_value)
        except Exception:
            num = 0.0

        if num == 0:
            return "background: rgba(74, 222, 128, 0.20); color: #dcfce7; font-weight: 700;"

        cap = 100000000.0
        ratio = min(abs(num), cap) / cap
        alpha = 0.12 + (0.36 * ratio)
        return f"background: rgba(248, 113, 113, {alpha:.3f}); color: #fecaca; font-weight: 700;"

    def _interpolate_hex(c1: str, c2: str, t: float) -> str:
        t = max(0.0, min(1.0, t))
        r1, g1, b1 = int(c1[0:2], 16), int(c1[2:4], 16), int(c1[4:6], 16)
        r2, g2, b2 = int(c2[0:2], 16), int(c2[2:4], 16), int(c2[4:6], 16)
        r = int(r1 + (r2 - r1) * t)
        g = int(g1 + (g2 - g1) * t)
        b = int(b1 + (b2 - b1) * t)
        return f"{r:02X}{g:02X}{b:02X}"

    def _excel_diff_fill(raw_value: float):
        if not OPENPYXL_AVAILABLE:
            return None
        try:
            num = float(raw_value)
        except Exception:
            num = 0.0

        if num == 0:
            color_hex = "DCFCE7"
        else:
            ratio = min(abs(num), 100000000.0) / 100000000.0
            color_hex = _interpolate_hex("FEE2E2", "B91C1C", ratio)
        return PatternFill(fill_type="solid", start_color=f"FF{color_hex}", end_color=f"FF{color_hex}")

    def build_detail_excel_bytes(table_data: list[list[str]], original_rows: list[dict]) -> bytes:
        if not OPENPYXL_AVAILABLE:
            return b""
        wb = Workbook()

        categories: dict[str, list[tuple[list[str], dict]]] = {
            "Excess Received": [],
            "Short Received": [],
            "Balanced": [],
        }

        for idx, row_vals in enumerate(table_data):
            src = original_rows[idx] if idx < len(original_rows) else {}
            diff_qty_raw = float(src.get("diff_qty") or 0)
            diff_val_raw = float(src.get("diff_val") or 0)
            if diff_qty_raw > 0 or (diff_qty_raw == 0 and diff_val_raw > 0):
                categories["Excess Received"].append((row_vals, src))
            elif diff_qty_raw < 0 or (diff_qty_raw == 0 and diff_val_raw < 0):
                categories["Short Received"].append((row_vals, src))
            else:
                categories["Balanced"].append((row_vals, src))

        def _write_sheet(ws, grouped_rows: list[tuple[list[str], dict]]):
            qty_number_fmt = "#,##0"
            val_number_fmt = "#,##0.00"

            for col_idx, col_name in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color="FFE2E8F0")
                cell.fill = PatternFill(fill_type="solid", start_color="FF1E3A8A", end_color="FF1E3A8A")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Add per-sheet totals row directly under the header.
            total_qty_loaded = sum(float(src.get("qty_loaded") or 0) for _, src in grouped_rows)
            total_value_loaded = sum(float(src.get("value_loaded") or 0) for _, src in grouped_rows)
            total_qty_offloaded = sum(float(src.get("qty_offloaded") or 0) for _, src in grouped_rows)
            total_value_offloaded = sum(float(src.get("value_offloaded") or 0) for _, src in grouped_rows)
            total_diff_qty = sum(float(src.get("diff_qty") or 0) for _, src in grouped_rows)
            total_diff_val = sum(float(src.get("diff_val") or 0) for _, src in grouped_rows)

            totals_row = [
                "TOTAL", "", "", "", "", "",
                total_qty_loaded,
                total_value_loaded,
                total_qty_offloaded,
                total_value_offloaded,
                total_diff_qty,
                total_diff_val,
            ]

            for col_idx, val in enumerate(totals_row, start=1):
                c = ws.cell(row=2, column=col_idx, value=val)
                c.font = Font(bold=True)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill = PatternFill(fill_type="solid", start_color="FFE5E7EB", end_color="FFE5E7EB")

            ws.cell(row=2, column=7).number_format = qty_number_fmt
            ws.cell(row=2, column=8).number_format = val_number_fmt
            ws.cell(row=2, column=9).number_format = qty_number_fmt
            ws.cell(row=2, column=10).number_format = val_number_fmt
            ws.cell(row=2, column=11).number_format = qty_number_fmt
            ws.cell(row=2, column=12).number_format = val_number_fmt

            ws.cell(row=2, column=11).fill = _excel_diff_fill(total_diff_qty)
            ws.cell(row=2, column=12).fill = _excel_diff_fill(total_diff_val)

            for row_num, (row_vals, src) in enumerate(grouped_rows, start=3):
                numeric_values = {
                    7: float(src.get("qty_loaded") or 0),
                    8: float(src.get("value_loaded") or 0),
                    9: float(src.get("qty_offloaded") or 0),
                    10: float(src.get("value_offloaded") or 0),
                    11: float(src.get("diff_qty") or 0),
                    12: float(src.get("diff_val") or 0),
                }
                for col_idx, val in enumerate(row_vals, start=1):
                    if col_idx in numeric_values:
                        val = numeric_values[col_idx]
                    c = ws.cell(row=row_num, column=col_idx, value=val)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    if col_idx in (7, 9, 11):
                        c.number_format = qty_number_fmt
                    elif col_idx in (8, 10, 12):
                        c.number_format = val_number_fmt

                qty_fill = _excel_diff_fill(float(src.get("diff_qty") or 0))
                val_fill = _excel_diff_fill(float(src.get("diff_val") or 0))
                ws.cell(row=row_num, column=11).fill = qty_fill
                ws.cell(row=row_num, column=12).fill = val_fill
                ws.cell(row=row_num, column=11).font = Font(bold=True)
                ws.cell(row=row_num, column=12).font = Font(bold=True)

            widths = [14, 28, 24, 14, 12, 42, 14, 16, 18, 18, 12, 14]
            for idx, w in enumerate(widths, start=1):
                ws.column_dimensions[chr(64 + idx)].width = w

        first = True
        for sheet_name in ["Excess Received", "Short Received", "Balanced"]:
            if first:
                ws = wb.active
                ws.title = sheet_name
                first = False
            else:
                ws = wb.create_sheet(title=sheet_name)
            _write_sheet(ws, categories[sheet_name])

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    cell_styles: dict[tuple[int, int], str] = {}

    for r in rows:
        shop_code = str(r.get("shop_code", "") or "").strip().upper()
        shop_meta = shop_meta_map.get(shop_code, {})
        manager_name = str(shop_meta.get("manager_name", "") or "").strip()
        shop_name = str(shop_meta.get("shop_name", "") or "").strip()
        date_text = r.get("date").strftime("%d %b %Y") if r.get("date") else ""
        diff_qty_raw = float(r.get("diff_qty") or 0)
        diff_val_raw = float(r.get("diff_val") or 0)
        row_display = [
            date_text,
            manager_name,
            shop_name,
            str(r.get("vehicle_no", "") or ""),
            str(r.get("item_code", "") or ""),
            str(r.get("item_name", "") or ""),
            num_fmt(r.get("qty_loaded", 0), 0),
            num_fmt(r.get("value_loaded", 0), 2),
            num_fmt(r.get("qty_offloaded", 0), 0),
            num_fmt(r.get("value_offloaded", 0), 2),
            num_fmt(diff_qty_raw, 0),
            num_fmt(diff_val_raw, 2),
        ]
        row_sort = [
            str(r.get("date") or ""),
            manager_name,
            shop_name,
            str(r.get("vehicle_no", "") or ""),
            str(r.get("item_code", "") or ""),
            str(r.get("item_name", "") or ""),
            str(float(r.get("qty_loaded") or 0)),
            str(float(r.get("value_loaded") or 0)),
            str(float(r.get("qty_offloaded") or 0)),
            str(float(r.get("value_offloaded") or 0)),
            str(diff_qty_raw),
            str(diff_val_raw),
        ]
        row_idx = len(table_rows)
        table_rows.append(row_display)
        sort_rows.append(row_sort)
        cell_styles[(row_idx, 10)] = diff_cell_style(diff_qty_raw)
        cell_styles[(row_idx, 11)] = diff_cell_style(diff_val_raw)

    detail_total_row = [
        "TOTAL", "", "", "", "", "",
        num_fmt(sum(float(r.get("qty_loaded") or 0) for r in rows), 0),
        num_fmt(sum(float(r.get("value_loaded") or 0) for r in rows), 2),
        num_fmt(sum(float(r.get("qty_offloaded") or 0) for r in rows), 0),
        num_fmt(sum(float(r.get("value_offloaded") or 0) for r in rows), 2),
        num_fmt(sum(float(r.get("diff_qty") or 0) for r in rows), 0),
        num_fmt(sum(float(r.get("diff_val") or 0) for r in rows), 2),
    ]

    excel_b64 = None
    download_name = "offloading_vs_loading_detail.csv"
    if OPENPYXL_AVAILABLE:
        export_source_rows = export_rows if export_rows is not None else rows
        export_table_rows = []
        for r in export_source_rows:
            shop_code = str(r.get("shop_code", "") or "").strip().upper()
            shop_meta = shop_meta_map.get(shop_code, {})
            manager_name = str(shop_meta.get("manager_name", "") or "").strip()
            shop_name = str(shop_meta.get("shop_name", "") or "").strip()
            date_text = r.get("date").strftime("%d %b %Y") if r.get("date") else ""
            diff_qty_raw = float(r.get("diff_qty") or 0)
            diff_val_raw = float(r.get("diff_val") or 0)
            export_table_rows.append([
                date_text,
                manager_name,
                shop_name,
                str(r.get("vehicle_no", "") or ""),
                str(r.get("item_code", "") or ""),
                str(r.get("item_name", "") or ""),
                num_fmt(r.get("qty_loaded", 0), 0),
                num_fmt(r.get("value_loaded", 0), 2),
                num_fmt(r.get("qty_offloaded", 0), 0),
                num_fmt(r.get("value_offloaded", 0), 2),
                num_fmt(diff_qty_raw, 0),
                num_fmt(diff_val_raw, 2),
            ])

        excel_bytes = build_detail_excel_bytes(export_table_rows, export_source_rows)
        excel_b64 = base64.b64encode(excel_bytes).decode("utf-8") if excel_bytes else None
        if excel_b64:
            download_name = "offloading_vs_loading_detail.xlsx"

    render_sortable_html_table(
        title="Loading vs Offloading (Detail)",
        headers=headers,
        rows=table_rows,
        sort_values=sort_rows,
        table_id="ovl-detail-table",
        min_width_px=1120,
        show_title=False,
        cell_styles=cell_styles,
        enable_download_hover=True,
        download_file_name=download_name,
        excel_base64=excel_b64,
        total_row=detail_total_row,
    )


def render_shop_summary_table(rows: list[dict], shop_meta_map: dict[str, dict[str, str]]) -> None:
    if not rows:
        st.info("No rows to display.")
        return

    def num_fmt(v, decimals=2):
        try:
            return f"{float(v):,.{decimals}f}"
        except Exception:
            return "0.00" if decimals == 2 else "0"

    table_rows = []
    sort_rows = []
    for r in rows:
        shop_code = str(r.get("shop_code", "") or "").strip().upper()
        shop_meta = shop_meta_map.get(shop_code, {})
        manager_name = str(shop_meta.get("manager_name", "") or "").strip()
        shop_name = str(shop_meta.get("shop_name", "") or "").strip()
        short_qty = float(r.get("short_qty") or 0)
        short_val = float(r.get("short_val") or 0)
        excess_qty = float(r.get("excess_qty") or 0)
        excess_val = float(r.get("excess_val") or 0)
        diff_qty = float(r.get("diff_qty") or 0)
        diff_val = float(r.get("diff_val") or 0)
        table_rows.append([
            manager_name,
            shop_name,
            num_fmt(short_qty, 0),
            num_fmt(short_val, 2),
            num_fmt(excess_qty, 0),
            num_fmt(excess_val, 2),
            num_fmt(diff_qty, 0),
            num_fmt(diff_val, 2),
        ])
        sort_rows.append([
            manager_name,
            shop_name,
            str(short_qty),
            str(short_val),
            str(excess_qty),
            str(excess_val),
            str(diff_qty),
            str(diff_val),
        ])

    shop_total_row = [
        "TOTAL",
        "",
        num_fmt(sum(float(r.get("short_qty") or 0) for r in rows), 0),
        num_fmt(sum(float(r.get("short_val") or 0) for r in rows), 2),
        num_fmt(sum(float(r.get("excess_qty") or 0) for r in rows), 0),
        num_fmt(sum(float(r.get("excess_val") or 0) for r in rows), 2),
        num_fmt(sum(float(r.get("diff_qty") or 0) for r in rows), 0),
        num_fmt(sum(float(r.get("diff_val") or 0) for r in rows), 2),
    ]

    render_sortable_html_table(
        title="Exception Shopwise Difference",
        headers=["Shop Manager", "Shop Name", "SHOP RECVD SHORT QTY", "SHOP RECVD SHORT VALUE", "SHOP RECVD EXCESS QTY", "SHOP RECVD EXCESS VALUE", "Final Difference QTY", "Final Difference Value"],
        rows=table_rows,
        sort_values=sort_rows,
        table_id="ovl-shop-table",
        min_width_px=1220,
        total_row=shop_total_row,
    )


def render_detail_filter_header(
    available_shop_codes: list[str],
    shop_meta_map: dict[str, dict[str, str]],
    detail_rows: list[dict],
    clicked_shop: str | None = None,
) -> tuple[str, str, list[str]]:
    pending_action = st.session_state.pop("detail_item_filter_action", None)
    if pending_action == "clear":
        # Apply search reset before widget instantiation to avoid Streamlit state mutation errors.
        st.session_state["detail_item_search"] = ""

    options = ["All"] + sorted(available_shop_codes)
    preferred = str(clicked_shop or st.session_state.get("detail_shop_filter", "All") or "All").strip().upper()
    if preferred not in options:
        preferred = "All"

    receive_options = ["All", "Short Received", "Excess Received"]
    receive_preferred = str(st.session_state.get("detail_receive_filter", "All") or "All")
    if receive_preferred not in receive_options:
        receive_preferred = "All"

    title_col, search_col, receive_col, filter_col = st.columns([0.44, 0.24, 0.14, 0.18], gap="small")
    with search_col:
        item_search = st.text_input(
            "Search Item Name",
            key="detail_item_search",
            placeholder="Type item name e.g. indomie",
            label_visibility="collapsed",
        ).strip()
    with receive_col:
        selected_receive = st.selectbox(
            "Receive Filter",
            options=receive_options,
            index=receive_options.index(receive_preferred),
            key="detail_receive_filter",
            label_visibility="collapsed",
        )
    with filter_col:
        selected = st.selectbox(
            "Detail Shop",
            options=options,
            index=options.index(preferred),
            key="detail_shop_filter",
            label_visibility="collapsed",
        )

    title_text = "Loading vs Offloading (Detail)"
    if selected != "All":
        meta = shop_meta_map.get(selected, {})
        shop_name = str(meta.get("shop_name", "") or "").strip()
        manager_name = str(meta.get("manager_name", "") or "").strip()
        if shop_name:
            title_text = f"Loading vs Offloading (Detail) — {selected} ({shop_name})"
        else:
            title_text = f"Loading vs Offloading (Detail) — {selected}"
        if manager_name:
            title_text = f"{title_text} | Manager: {manager_name}"
    with title_col:
        st.markdown(f"<div class='section-title'>{html.escape(title_text)}</div>", unsafe_allow_html=True)

    scoped_rows = detail_rows if selected == "All" else [
        row for row in detail_rows
        if str(row.get("shop_code", "") or "").strip().upper() == selected
    ]

    if selected_receive == "Short Received":
        scoped_rows = [row for row in scoped_rows if float(row.get("diff_val") or 0) < 0]
    elif selected_receive == "Excess Received":
        scoped_rows = [row for row in scoped_rows if float(row.get("diff_val") or 0) > 0]

    all_item_names = sorted({
        str(row.get("item_name", "") or "").strip()
        for row in scoped_rows
        if str(row.get("item_name", "") or "").strip()
    })
    if item_search:
        matching_item_names = [name for name in all_item_names if item_search.lower() in name.lower()]
    else:
        matching_item_names = all_item_names

    if pending_action == "select_all":
        st.session_state["detail_item_filter"] = matching_item_names
    elif pending_action == "clear":
        st.session_state["detail_item_filter"] = []

    existing_selection = st.session_state.get("detail_item_filter", []) or []
    valid_selection = [name for name in existing_selection if name in matching_item_names]
    if valid_selection != existing_selection:
        st.session_state["detail_item_filter"] = valid_selection

    selected_items: list[str] = []
    if item_search or valid_selection:
        picker_col, select_all_col, clear_col = st.columns([0.76, 0.12, 0.12], gap="small")
        with picker_col:
            selected_items = st.multiselect(
                "Matching Item Names",
                options=matching_item_names,
                key="detail_item_filter",
                placeholder="Select one or more matching item names",
                help="Type part of an item name above to narrow the list, then select one or more matches.",
            )
        with select_all_col:
            st.markdown('<div style="margin-top: 1.75rem;"></div>', unsafe_allow_html=True)
            if st.button("Select All", key="detail_item_filter_all", use_container_width=True, disabled=not matching_item_names):
                st.session_state["detail_item_filter_action"] = "select_all"
                st.rerun()
        with clear_col:
            st.markdown('<div style="margin-top: 1.75rem;"></div>', unsafe_allow_html=True)
            if st.button("Clear", key="detail_item_filter_clear", use_container_width=True, disabled=not (item_search or selected_items)):
                st.session_state["detail_item_filter_action"] = "clear"
                st.rerun()

        if item_search:
            st.caption(f"{len(matching_item_names):,} item name(s) matched '{item_search}'.")

    return selected, selected_receive, selected_items


def render_header_with_date_filter(min_date, max_date):
    def _coerce_date(val):
        if isinstance(val, date):
            return val
        if isinstance(val, datetime):
            return val.date()
        if isinstance(val, str):
            try:
                return datetime.fromisoformat(val).date()
            except Exception:
                return None
        return None

    def _clamp_date(val, lo, hi):
        if val is None:
            return None
        if lo is not None and val < lo:
            return lo
        if hi is not None and val > hi:
            return hi
        return val

    yesterday = (datetime.today() - timedelta(days=1)).date()
    if max_date and max_date <= yesterday:
        default_end = max_date
    else:
        default_end = yesterday
    default_end = _clamp_date(default_end, min_date, max_date)
    default_start = min_date
    default_start = _clamp_date(default_start, min_date, max_date)

    if default_start is None and default_end is not None:
        default_start = default_end
    if default_end is None and default_start is not None:
        default_end = default_start

    saved_raw = st.session_state.get("offload_date_range", (default_start, default_end))
    if isinstance(saved_raw, (tuple, list)):
        parsed = [_coerce_date(v) for v in saved_raw]
        parsed = [p for p in parsed if p is not None]
        if len(parsed) >= 2:
            saved = (parsed[0], parsed[1])
        elif len(parsed) == 1:
            saved = (parsed[0], parsed[0])
        else:
            saved = (default_start, default_end)
    else:
        parsed_one = _coerce_date(saved_raw)
        saved = (parsed_one, parsed_one) if parsed_one else (default_start, default_end)

    saved = (
        _clamp_date(saved[0], min_date, max_date),
        _clamp_date(saved[1], min_date, max_date),
    )
    if saved[0] is None and saved[1] is not None:
        saved = (saved[1], saved[1])
    elif saved[1] is None and saved[0] is not None:
        saved = (saved[0], saved[0])
    elif saved[0] is None and saved[1] is None:
        saved = (default_start, default_end)

    if saved[0] > saved[1]:
        saved = (saved[1], saved[0])

    is_master = st.session_state.get("ovl_is_master_user", False)
    if is_master and WMS_SYNC_AVAILABLE:
        hdr_col, date_col, update_col, refresh_col, logout_col = st.columns([0.48, 0.20, 0.13, 0.08, 0.08], gap="small")
    else:
        update_col = None
        hdr_col, date_col, refresh_col, logout_col = st.columns([0.62, 0.20, 0.09, 0.09])

    with hdr_col:
        st.markdown(
            """
            <div style="display:flex;align-items:center;gap:12px;padding:4px 0 8px 0;">
                <img src="https://melcom.com/media/favicon/stores/1/faviconn_162_x_184px.jpg"
                     style="height:34px;border-radius:50%;border:2px solid #3b82f6;">
                <div>
                    <div class="dashboard-title">Loading vs Offloading Dashboard</div>
                    <div class="dashboard-sub">WH Database · Tabular reconciliation view</div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with date_col:
        st.markdown(
            '<div style="text-align:right;font-size:0.72rem;color:#6ee7b7;font-weight:600;margin-bottom:2px;">Date Range</div>',
            unsafe_allow_html=True,
        )
        selection = st.date_input(
            "offload_date_pick",
            value=(saved[0], saved[1]),
            min_value=min_date,
            max_value=max_date,
            key="offload_date_pick",
            label_visibility="collapsed",
        )

    today     = datetime.today().date()
    yesterday = today - timedelta(days=1)

    # ── Monday: also upload Saturday (yesterday-1) ───────────────────────────
    # yesterday = Sunday on Monday; Saturday is two days ago
    if today.weekday() == 0:  # 0 = Monday
        saturday = yesterday - timedelta(days=1)
        _dates_to_upload = [saturday, yesterday]   # Saturday then Sunday
    else:
        _dates_to_upload = [yesterday]

    # ── Show any pending email notification that survived st.rerun() ──────────
    _pending = st.session_state.pop("_lvo_email_notify", None)
    if _pending:
        _kind, _msg = _pending
        if _kind == "success":
            st.success(_msg)
        else:
            st.warning(_msg)

    if update_col is not None:
        with update_col:
            st.markdown('<div style="margin-top:18px;"></div>', unsafe_allow_html=True)
            already_loaded = all(_wms_sync.date_exists_in_db(d) for d in _dates_to_upload)
            if today.weekday() == 0:
                _date_label_btn = f"Sat {_dates_to_upload[0].strftime('%d %b')} + Sun {_dates_to_upload[1].strftime('%d %b')}"
            else:
                _date_label_btn = yesterday.strftime('%d %b %Y')
            btn_label = "✅ Updated" if already_loaded else "⬆ Update Data"
            btn_help  = (
                f"Data for {_date_label_btn} is already loaded. Click to force re-import."
                if already_loaded else
                f"Load {_date_label_btn} data from network share and send email"
            )
            if st.button(btn_label, key="wms_update_btn", use_container_width=True, help=btn_help):
                st.session_state["wms_sync_trigger"] = True
                st.session_state["wms_sync_force"] = already_loaded
                st.rerun()

        # ── File-missing banner (shown below header, full width) ──────────────
        if not already_loaded:
            _file_status = _wms_check_file_status(yesterday)
            if _file_status == "network_error":
                st.markdown(
                    f"""
                    <div style="background:rgba(248,81,73,.12);border:1px solid rgba(248,81,73,.4);
                                border-radius:10px;padding:10px 16px;margin:4px 0 8px 0;
                                display:flex;align-items:center;gap:10px;">
                      <span style="font-size:18px;">🔌</span>
                      <span style="color:#f85149;font-weight:600;font-size:13px;">
                        Network share <code>\\\\10.10.0.30\\mis</code> is not accessible.
                        Check VPN / network connectivity.
                      </span>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
            elif _file_status == "missing":
                _expected = f"WMS_shopdiff_{yesterday.strftime('%d_%m_%Y')}.csv"
                st.markdown(
                    f"""
                    <div style="background:rgba(255,193,7,.10);border:1px solid rgba(255,193,7,.4);
                                border-radius:10px;padding:10px 16px;margin:4px 0 8px 0;
                                display:flex;align-items:center;gap:10px;">
                      <span style="font-size:18px;">⚠️</span>
                      <span style="color:#ffc107;font-weight:600;font-size:13px;">
                        File not yet available on network share —
                        <code>{_expected}</code> not found in
                        <code>\\\\10.10.0.30\\mis</code>.
                        Dashboard data may be incomplete for {yesterday.strftime('%d %b %Y')}.
                      </span>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

    with refresh_col:
        st.markdown('<div style="margin-top:18px;"></div>', unsafe_allow_html=True)
        if st.button("🔄", key="offload_refresh", help="Clear cache and reload"):
            st.cache_data.clear()
            st.rerun()

    with logout_col:
        st.markdown('<div style="margin-top:18px;"></div>', unsafe_allow_html=True)
        if st.button("Logout", key="header_logout_btn", use_container_width=True, help="Sign out"):
            st.session_state["ovl_authenticated"] = False
            st.session_state["ovl_user"] = None
            st.session_state["ovl_is_master_user"] = False
            st.session_state["ovl_allowed_shops"] = []
            st.rerun()

    # ── CTA: Update Data — same pipeline as home_dashboard LVO_offloading_vs_loading ──
    if st.session_state.pop("wms_sync_trigger", False) and WMS_SYNC_AVAILABLE:
        force = st.session_state.pop("wms_sync_force", False)

        if not force and all(_wms_sync.date_exists_in_db(d) for d in _dates_to_upload):
            st.info(
                f"Data for {_date_label_btn} is already loaded. "
                "Click **⬆ Update Data** again to force re-import."
            )
        else:
            _status_title = (
                f"Updating data for {_date_label_btn}…"
            )
            with st.status(_status_title, expanded=True) as _sync_status:
                try:
                    _total_rows   = 0
                    _loaded_dates = []
                    _last_csv_path = None

                    _col_sql      = ", ".join(f'"{c}"' for c in _LVO_COLUMNS)
                    _null         = lambda v: None if pd.isnull(v) else v
                    _hist_col_sql = ", ".join(
                        f'"{c}"' for c in _LVO_COLUMNS + ["history_action", "history_at", "source_file"]
                    )

                    # ── Ensure table structure once ───────────────────────────
                    _conn_init = psycopg2.connect(**DB_CONFIG)
                    try:
                        _wms_sync._ensure_loadingvsoffloading_tables(_conn_init)
                        _conn_init.commit()
                    finally:
                        _conn_init.close()

                    # ── Loop over each date to upload ─────────────────────────
                    for _target_date in _dates_to_upload:
                        _d_label = _target_date.strftime('%d %b %Y')

                        # Step 1: find CSV
                        _sync_status.write(
                            f"🔍 [{_d_label}] Searching for "
                            f"`WMS_shopdiff_{_target_date.strftime('%d_%m_%Y')}.csv`…"
                        )
                        try:
                            _csv_path = _wms_sync.find_wms_file(_target_date)
                        except OSError as _e:
                            _sync_status.update(label="❌ Network share not accessible", state="error")
                            st.error(str(_e))
                            raise

                        if not _csv_path:
                            _sync_status.write(
                                f"⚠️ [{_d_label}] File not found — skipping this date."
                            )
                            continue

                        _sync_status.write(f"📄 [{_d_label}] Found: `{os.path.basename(_csv_path)}`")
                        _last_csv_path = _csv_path

                        # Step 2: read CSV
                        _sync_status.write(f"📊 [{_d_label}] Reading CSV…")
                        _df_raw = pd.read_csv(_csv_path, encoding="utf-8-sig", dtype=str)

                        # Step 3: normalise
                        _df_norm = _wms_sync._prepare_lvo_offloading_df(
                            _df_raw, os.path.basename(_csv_path)
                        )
                        _df_norm = _wms_sync._rename_by_expected_columns(_df_norm, _LVO_COLUMNS)
                        _missing = [c for c in _LVO_COLUMNS if c not in _df_norm.columns]
                        if _missing:
                            raise ValueError(f"Missing columns after normalisation: {', '.join(_missing)}")
                        _df_up = _df_norm[_LVO_COLUMNS].copy()
                        _df_up["date"] = pd.to_datetime(_df_up["date"], errors="coerce").dt.date
                        for _nc in ["qty_loaded", "value_loaded", "qty_offloaded",
                                    "value_offloaded", "diff_qty", "diff_val"]:
                            _df_up[_nc] = pd.to_numeric(_df_up[_nc], errors="coerce").fillna(0)
                        _df_up = _df_up[_df_up["date"] == _target_date].copy()
                        if _df_up.empty:
                            _sync_status.write(
                                f"⚠️ [{_d_label}] No rows found after normalisation — skipping."
                            )
                            continue
                        _sync_status.write(
                            f"✅ [{_d_label}] Normalised **{len(_df_up):,} rows**"
                        )

                        # Step 4: DB operations
                        _conn = psycopg2.connect(**DB_CONFIG)
                        _conn.autocommit = False
                        try:
                            with _conn.cursor() as _cur:
                                _cur.execute(
                                    "DELETE FROM public.offloading_vs_loading WHERE date = %s",
                                    (_target_date,),
                                )
                                _deleted = _cur.rowcount
                            if _deleted:
                                _sync_status.write(
                                    f"🗑️ [{_d_label}] Removed **{_deleted}** stale rows"
                                )

                            _values = [
                                tuple(_null(row.get(c)) for c in _LVO_COLUMNS)
                                for _, row in _df_up.iterrows()
                            ]
                            with _conn.cursor() as _cur:
                                execute_values(
                                    _cur,
                                    f"INSERT INTO public.offloading_vs_loading ({_col_sql}) VALUES %s",
                                    _values,
                                    page_size=5000,
                                )
                                _cur.execute(
                                    "UPDATE public.offloading_vs_loading "
                                    "SET diff = diff_qty WHERE diff IS DISTINCT FROM diff_qty"
                                )

                            _now_ts = datetime.now()
                            _hist_values = [
                                tuple(_null(row.get(c)) for c in _LVO_COLUMNS)
                                + ("auto_sync", _now_ts, os.path.basename(_csv_path))
                                for _, row in _df_up.iterrows()
                            ]
                            with _conn.cursor() as _cur:
                                execute_values(
                                    _cur,
                                    f"INSERT INTO public.offloading_vs_loading_history ({_hist_col_sql}) VALUES %s",
                                    _hist_values,
                                    page_size=5000,
                                )
                                _cur.execute(
                                    "UPDATE public.offloading_vs_loading_history "
                                    "SET diff = diff_qty "
                                    "WHERE history_action = 'auto_sync' "
                                    "AND history_at >= NOW() - INTERVAL '10 minutes'"
                                )

                            _conn.commit()
                            _sync_status.write(
                                f"✅ [{_d_label}] Committed **{len(_values):,} rows**"
                            )
                            _total_rows += len(_values)
                            _loaded_dates.append(_target_date)

                        except Exception:
                            _conn.rollback()
                            raise
                        finally:
                            _conn.close()

                    # ── Summary label ─────────────────────────────────────────
                    _dates_summary = " + ".join(d.strftime('%d %b %Y') for d in _loaded_dates)
                    _sync_status.update(
                        label=f"✅ {_total_rows:,} rows loaded for {_dates_summary or _date_label_btn}",
                        state="complete",
                    )

                    # ── Email — send once for the latest loaded date ──────────
                    if _loaded_dates:
                        _email_date = _loaded_dates[-1]   # most recent (Sunday on Monday)
                        _sync_status.write("📧 Sending email notification…")
                        try:
                            _excel = _wms_sync.build_excel_bytes(_email_date)
                            _wms_sync.send_email(_email_date, _total_rows, _excel)
                            _sync_status.write("📧 Email sent ✓")
                            st.session_state["_lvo_email_notify"] = (
                                "success",
                                f"✅ Data loaded & email sent — "
                                f"**{_total_rows:,} rows** for {_dates_summary}",
                            )
                        except Exception as _email_err:
                            _sync_status.write(f"⚠️ Email error: {_email_err}")
                            st.session_state["_lvo_email_notify"] = (
                                "warning",
                                f"✅ Data loaded ({_total_rows:,} rows for {_dates_summary}) but "
                                f"**email failed** — {_email_err}",
                            )

                        st.cache_data.clear()
                        st.rerun()

                except Exception as _err:
                    _sync_status.update(label="❌ Upload failed", state="error")
                    st.error(f"❌ {_err}")

    if isinstance(selection, tuple) and len(selection) == 2:
        d_start = _coerce_date(selection[0])
        d_end = _coerce_date(selection[1])
        if d_start is None and d_end is None:
            d_start, d_end = default_start, default_end
        elif d_start is None:
            d_start = d_end
        elif d_end is None:
            d_end = d_start
    elif isinstance(selection, (tuple, list)) and len(selection) == 1:
        only_date = _coerce_date(selection[0]) or default_end
        d_start = d_end = only_date
    else:
        single = _coerce_date(selection) or default_end
        d_start = d_end = single

    if d_start > d_end:
        d_start, d_end = d_end, d_start

    d_start = _clamp_date(d_start, min_date, max_date)
    d_end = _clamp_date(d_end, min_date, max_date)

    st.session_state["offload_date_range"] = (d_start, d_end)
    return d_start, d_end


def main() -> None:
    inject_css()
    init_auth_state()

    if not st.session_state.get("ovl_authenticated", False):
        render_login_gate()
        return

    if render_logged_in_sidebar():
        return

    user = st.session_state.get("ovl_user") or {}
    _ta = str((user or {}).get("table_access", "") or "")
    _is_super = "all" in {x.strip().lower() for x in _ta.split(",") if x.strip()}
    if _is_super:
        is_master_user = True
        allowed_shops: list[str] = []
    else:
        is_master_user, allowed_shops = get_user_shop_scope(
            str(user.get("employee_id", "") or ""),
            str(user.get("full_name", "") or ""),
        )
    st.session_state["ovl_is_master_user"] = is_master_user
    st.session_state["ovl_allowed_shops"] = allowed_shops

    if not _is_super and not is_master_user and not allowed_shops:
        st.error("No shop mapping found in latest century_penetration.opsmgr upload for this user.")
        return

    latest_upload_date = get_latest_opsmgr_upload_date()

    min_date, max_date = load_date_bounds()
    if min_date is None or max_date is None:
        live_count, live_min, live_max = get_live_date_bounds_uncached()
        if live_count > 0 and live_min is not None and live_max is not None:
            # One-time self-heal for stale Streamlit cache after uploads/restores.
            if not st.session_state.get("offload_bounds_cache_refreshed", False):
                st.session_state["offload_bounds_cache_refreshed"] = True
                st.cache_data.clear()
                st.rerun()
            min_date, max_date = live_min, live_max
        else:
            st.error("No data found in `offloading_vs_loading`. Run setup script first.")
            return
    else:
        st.session_state["offload_bounds_cache_refreshed"] = False

    start_date, end_date = render_header_with_date_filter(min_date, max_date)

    data = load_offloading_vs_loading(start_date, end_date)
    data = apply_shop_scope(data, allowed_shops, is_master_user)
    if not data:
        st.info("No records for selected date range and access scope.")
        return

    render_metrics(data, start_date, end_date)
    shop_rows = load_shop_summary(start_date, end_date)
    shop_rows = apply_shop_scope(shop_rows, allowed_shops, is_master_user)
    shop_meta_map = load_shop_meta_map()

    if is_master_user:
        if latest_upload_date:
            st.caption(f"Access Scope: ALL shops (master user) | opsmgr upload_date: {latest_upload_date}")
        else:
            st.caption("Access Scope: ALL shops (master user)")
    else:
        if latest_upload_date:
            st.caption(f"Access Scope: {len(allowed_shops)} mapped shop(s) from latest opsmgr upload_date: {latest_upload_date}")
        else:
            st.caption(f"Access Scope: {len(allowed_shops)} mapped shop(s) from century_penetration.opsmgr")

    render_shop_summary_table(shop_rows, shop_meta_map)

    available_shop_codes = sorted({str(r.get("shop_code", "") or "").strip().upper() for r in data if str(r.get("shop_code", "") or "").strip()})
    selected_shop, selected_receive, selected_items = render_detail_filter_header(available_shop_codes, shop_meta_map, data)
    filtered_detail = data if selected_shop == "All" else [
        r for r in data
        if str(r.get("shop_code", "") or "").strip().upper() == selected_shop
    ]

    if selected_receive == "Short Received":
        filtered_detail = [r for r in filtered_detail if float(r.get("diff_val") or 0) < 0]
    elif selected_receive == "Excess Received":
        filtered_detail = [r for r in filtered_detail if float(r.get("diff_val") or 0) > 0]

    if selected_items:
        selected_item_names = set(selected_items)
        filtered_detail = [
            r for r in filtered_detail
            if str(r.get("item_name", "") or "").strip() in selected_item_names
        ]

    render_detail_table(filtered_detail, shop_meta_map, export_rows=data)


if __name__ == "__main__":
    main()
