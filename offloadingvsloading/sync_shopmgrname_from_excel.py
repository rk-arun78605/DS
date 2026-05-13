import os
import re
from datetime import datetime
from typing import Iterable

import psycopg2
from psycopg2.extras import execute_values
from openpyxl import load_workbook

DB_CONFIG = {
    "host": "localhost",
    "port": 3307,
    "user": "postgres",
    "password": "hello",
    "dbname": "WH",
}

EXCEL_PATH = os.path.join(os.environ.get("USERPROFILE", ""), "Downloads", "shop managerlist.xlsx")


def _clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _clean_manager_name(value) -> str:
    txt = _clean_text(value)
    txt = re.sub(r"\s*\(SHOP MANAGER\)\s*", " ", txt, flags=re.IGNORECASE)
    txt = re.sub(r"\s{2,}", " ", txt)
    return txt.strip()


def read_excel_rows(path: str) -> list[tuple[str, str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    header_map = {str(h).strip().lower(): idx for idx, h in enumerate(header) if h is not None}

    code_idx = header_map.get("shop_code")
    desc_idx = header_map.get("shop name")
    mgr_idx = header_map.get("manager name")

    if code_idx is None or desc_idx is None or mgr_idx is None:
        raise ValueError("Excel must have columns: Shop_Code, Shop Name, Manager Name")

    rows: set[tuple[str, str, str]] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        shop_code = _clean_text(row[code_idx]).upper()
        shop_description = _clean_text(row[desc_idx])
        shop_manager_name = _clean_manager_name(row[mgr_idx])
        if not shop_code or not shop_manager_name:
            continue
        rows.add((shop_code, shop_description, shop_manager_name))

    return sorted(rows)


def ensure_table(conn) -> None:
    ddl = """
    CREATE TABLE IF NOT EXISTS public.shopmgrname (
        id BIGSERIAL PRIMARY KEY,
        shop_code TEXT NOT NULL,
        shop_description TEXT,
        shop_manager_name TEXT NOT NULL,
        valid_from TIMESTAMP NOT NULL DEFAULT NOW(),
        valid_to TIMESTAMP NULL,
        is_current BOOLEAN NOT NULL DEFAULT TRUE,
        source_file TEXT,
        load_batch_ts TIMESTAMP NOT NULL DEFAULT NOW(),
        updated_at TIMESTAMP NOT NULL DEFAULT NOW()
    );

    CREATE INDEX IF NOT EXISTS idx_shopmgrname_current_code
        ON public.shopmgrname (shop_code)
        WHERE is_current = TRUE;

    CREATE INDEX IF NOT EXISTS idx_shopmgrname_history_code_from
        ON public.shopmgrname (shop_code, valid_from DESC);

    CREATE UNIQUE INDEX IF NOT EXISTS uq_shopmgrname_current_triplet
        ON public.shopmgrname (shop_code, COALESCE(shop_description, ''), shop_manager_name)
        WHERE is_current = TRUE;
    """
    with conn.cursor() as cur:
        cur.execute(ddl)


def sync_rows(conn, rows: Iterable[tuple[str, str, str]], source_file: str) -> tuple[int, int]:
    ts = datetime.now()
    with conn.cursor() as cur:
        cur.execute("DROP TABLE IF EXISTS tmp_shopmgr_upload")
        cur.execute(
            """
            CREATE TEMP TABLE tmp_shopmgr_upload (
                shop_code TEXT NOT NULL,
                shop_description TEXT,
                shop_manager_name TEXT NOT NULL
            ) ON COMMIT DROP
            """
        )

        execute_values(
            cur,
            "INSERT INTO tmp_shopmgr_upload (shop_code, shop_description, shop_manager_name) VALUES %s",
            list(rows),
            page_size=1000,
        )

        cur.execute(
            """
            UPDATE public.shopmgrname t
            SET is_current = FALSE,
                valid_to = %(ts)s,
                updated_at = %(ts)s
            WHERE t.is_current = TRUE
              AND NOT EXISTS (
                    SELECT 1
                    FROM tmp_shopmgr_upload u
                    WHERE u.shop_code = t.shop_code
                      AND COALESCE(u.shop_description, '') = COALESCE(t.shop_description, '')
                      AND u.shop_manager_name = t.shop_manager_name
              )
            """,
            {"ts": ts},
        )
        deactivated = cur.rowcount

        cur.execute(
            """
            INSERT INTO public.shopmgrname (
                shop_code,
                shop_description,
                shop_manager_name,
                valid_from,
                valid_to,
                is_current,
                source_file,
                load_batch_ts,
                updated_at
            )
            SELECT
                u.shop_code,
                u.shop_description,
                u.shop_manager_name,
                %(ts)s,
                NULL,
                TRUE,
                %(src)s,
                %(ts)s,
                %(ts)s
            FROM tmp_shopmgr_upload u
            WHERE NOT EXISTS (
                SELECT 1
                FROM public.shopmgrname t
                WHERE t.is_current = TRUE
                  AND t.shop_code = u.shop_code
                  AND COALESCE(t.shop_description, '') = COALESCE(u.shop_description, '')
                  AND t.shop_manager_name = u.shop_manager_name
            )
            """,
            {"ts": ts, "src": source_file},
        )
        inserted = cur.rowcount

    return inserted, deactivated


def main() -> None:
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Excel not found: {EXCEL_PATH}")

    rows = read_excel_rows(EXCEL_PATH)
    if not rows:
        raise RuntimeError("No valid rows found in Excel file")

    with psycopg2.connect(**DB_CONFIG) as conn:
        ensure_table(conn)
        inserted, deactivated = sync_rows(conn, rows, EXCEL_PATH)
        conn.commit()

    print(f"Synced {len(rows)} unique rows from Excel")
    print(f"Inserted new current rows: {inserted}")
    print(f"Closed old rows (history): {deactivated}")


if __name__ == "__main__":
    main()
